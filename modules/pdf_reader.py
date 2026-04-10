import os
import re
import fitz  # PyMuPDF

# =============================================
# 📄 Document Reader Module
# Supports: PDF, DOCX, TXT, MD, PPTX, EPUB
# =============================================

SUPPORTED_EXTENSIONS = [".pdf", ".docx", ".txt", ".md", ".pptx", ".epub"]


def _has_meaningful_text(text: str, threshold: int = 24) -> bool:
    return len(re.sub(r"\s+", "", text or "")) >= threshold


def _ocr_page(page) -> str:
    try:
        import io
        from PIL import Image
        import pytesseract
    except ImportError:
        return ""

    try:
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        text = pytesseract.image_to_string(image)
        return text.strip()
    except pytesseract.pytesseract.TesseractNotFoundError:
        return ""
    except Exception:
        return ""


def _missing_ocr_dependencies() -> bool:
    try:
        import pytesseract
        from PIL import Image  # noqa: F401
        _ = pytesseract.get_tesseract_version()
        return False
    except Exception:
        return True


def read_file(file_path: str) -> str:
    """
    Auto-detects file type and extracts plain text.
    Supports .pdf, .docx, .txt, .md, .pptx, .epub
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return read_pdf(file_path)
    elif ext == ".docx":
        return read_docx(file_path)
    elif ext == ".txt":
        return read_txt(file_path)
    elif ext == ".md":
        return read_md(file_path)
    elif ext == ".pptx":
        return read_pptx(file_path)
    elif ext == ".epub":
        return read_epub(file_path)
    else:
        return f"❌ Unsupported file type: {ext}. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"


# -----------------------------------------------
# Existing readers (unchanged)
# -----------------------------------------------

def read_pdf(file_path: str) -> str:
    """Extracts text from a PDF file, page by page."""
    try:
        doc = fitz.open(file_path)
        pages = []
        used_ocr = False
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if not _has_meaningful_text(text):
                ocr_text = _ocr_page(page)
                if ocr_text:
                    text = ocr_text
                    used_ocr = True
            if text.strip():
                pages.append(f"--- Page {page_num} ---\n{text.strip()}")
        doc.close()
        full_text = "\n".join(pages).strip()
        if full_text:
            return full_text
        if _missing_ocr_dependencies():
            return (
                "⚠️ This PDF appears to be scanned or image-based. OCR support is not available. "
                "Install Pillow, pytesseract, and the Tesseract OCR app to extract scanned PDFs."
            )
        if used_ocr:
            return full_text
        return "⚠️ This PDF appears to be empty or scanned. Text extraction not possible."
    except Exception as e:
        return f"❌ Error reading PDF: {str(e)}"


def read_docx(file_path: str) -> str:
    """Extracts text from a DOCX file, paragraph by paragraph."""
    try:
        from docx import Document
        doc = Document(file_path)
        full_text = ""
        for para in doc.paragraphs:
            if para.text.strip():
                full_text += f"{para.text}\n"
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        full_text += f"{cell.text}\n"
        if not full_text.strip():
            return "⚠️ This DOCX appears to be empty. Text extraction not possible."
        return full_text.strip()
    except ImportError:
        return "❌ python-docx not installed. Run: pip install python-docx"
    except Exception as e:
        return f"❌ Error reading DOCX: {str(e)}"


# -----------------------------------------------
# New readers
# -----------------------------------------------

def read_txt(file_path: str) -> str:
    """Reads a plain .txt file with automatic encoding fallback."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                text = f.read()
            if text.strip():
                return text.strip()
        except (UnicodeDecodeError, LookupError):
            continue
        except Exception as e:
            return f"❌ Error reading TXT: {str(e)}"
    return "⚠️ Could not decode this text file."


def read_md(file_path: str) -> str:
    """
    Reads a Markdown file and strips formatting markers so the
    LM sees clean prose instead of raw # / ** / __ symbols.
    """
    raw = None
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                raw = f.read()
            break
        except (UnicodeDecodeError, LookupError):
            continue
        except Exception as e:
            return f"❌ Error reading Markdown: {str(e)}"

    if raw is None:
        return "⚠️ Could not decode this Markdown file."

    # Strip fenced code blocks — keep the code, drop the ``` fences
    raw = re.sub(r'```[\w]*\n?', '', raw)
    # Strip inline code backticks
    raw = re.sub(r'`([^`]+)`', r'\1', raw)
    # Strip heading markers (# ## ###)
    raw = re.sub(r'^#{1,6}\s+', '', raw, flags=re.MULTILINE)
    # Strip bold/italic markers
    raw = re.sub(r'(\*{1,3}|_{1,3})(.+?)\1', r'\2', raw)
    # Strip HTML tags (sometimes present in .md)
    raw = re.sub(r'<[^>]+>', '', raw)
    # Strip horizontal rules (--- / *** / ___)
    raw = re.sub(r'^[-*_]{3,}\s*$', '', raw, flags=re.MULTILINE)
    # Collapse excessive blank lines
    raw = re.sub(r'\n{3,}', '\n\n', raw)

    text = raw.strip()
    if not text:
        return "⚠️ This Markdown file appears to be empty."
    return text


def read_pptx(file_path: str) -> str:
    """
    Extracts text from a PowerPoint (.pptx) file slide by slide.
    Captures: title, body text, tables, and speaker notes.
    The model only ever sees plain text — it has no idea it came from slides.
    """
    try:
        from pptx import Presentation
    except ImportError:
        return "❌ python-pptx not installed. Run: pip install python-pptx"

    try:
        prs = Presentation(file_path)
        slides_text = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            parts = [f"[Slide {slide_num}]"]

            for shape in slide.shapes:
                # Text frames (titles, body boxes)
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = para.text.strip()
                        if line:
                            parts.append(line)
                # Tables
                elif shape.has_table:
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            parts.append(" | ".join(cells))

            # Speaker notes — often the richest content
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"[Notes] {notes}")

            if len(parts) > 1:   # more than just the [Slide N] header
                slides_text.append("\n".join(parts))

        if not slides_text:
            return "⚠️ This presentation appears to be empty or contains only images."

        return "\n\n".join(slides_text)
    except Exception as e:
        return f"❌ Error reading PPTX: {str(e)}"


def read_epub(file_path: str) -> str:
    """
    Extracts text from an EPUB ebook chapter by chapter.
    Requires: ebooklib + beautifulsoup4
    """
    try:
        import ebooklib
        from ebooklib import epub
        from bs4 import BeautifulSoup
    except ImportError:
        return "❌ ebooklib / beautifulsoup4 not installed. Run: pip install ebooklib beautifulsoup4"

    try:
        book = epub.read_epub(file_path, options={"ignore_ncx": True})
        chapters = []

        for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            raw_html = item.get_content().decode("utf-8", errors="replace")
            soup = BeautifulSoup(raw_html, "html.parser")

            # Remove nav/toc elements that are pure link lists
            for tag in soup.find_all(["nav", "aside"]):
                tag.decompose()

            text = soup.get_text(separator="\n")
            text = re.sub(r'\n{3,}', '\n\n', text).strip()
            if text:
                chapters.append(text)

        if not chapters:
            return "⚠️ This EPUB appears to be empty or DRM-protected."

        return "\n\n---\n\n".join(chapters)
    except Exception as e:
        return f"❌ Error reading EPUB: {str(e)}"


# -----------------------------------------------
# Metadata helpers
# -----------------------------------------------

def get_page_count(file_path: str) -> int:
    """
    Returns a meaningful unit count per file type:
    PDF  → pages       DOCX → paragraphs
    PPTX → slides      EPUB → chapters
    TXT/MD → estimated pages (~300 words each)
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        try:
            doc = fitz.open(file_path)
            count = len(doc)
            doc.close()
            return count
        except Exception:
            return 0
    elif ext == ".docx":
        try:
            from docx import Document
            doc = Document(file_path)
            return sum(1 for p in doc.paragraphs if p.text.strip())
        except Exception:
            return 0
    elif ext == ".pptx":
        try:
            from pptx import Presentation
            return len(Presentation(file_path).slides)
        except Exception:
            return 0
    elif ext == ".epub":
        try:
            import ebooklib
            from ebooklib import epub
            book = epub.read_epub(file_path, options={"ignore_ncx": True})
            return len(list(book.get_items_of_type(ebooklib.ITEM_DOCUMENT)))
        except Exception:
            return 0
    elif ext in (".txt", ".md"):
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                words = len(f.read().split())
            return max(1, round(words / 300))   # ~300 words per page
        except Exception:
            return 0
    return 0


def get_page_label(file_path: str) -> str:
    """Human-readable unit label for get_page_count."""
    ext = os.path.splitext(file_path)[1].lower()
    labels = {
        ".pdf":  "pages",
        ".docx": "paragraphs",
        ".pptx": "slides",
        ".epub": "chapters",
        ".txt":  "est. pages",
        ".md":   "est. pages",
    }
    return labels.get(ext, "pages")


# -----------------------------------------------
# Chunker (unchanged)
# -----------------------------------------------

def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list:
    """
    Splits large text into smaller overlapping chunks.
    chunk_size = words per chunk
    overlap    = shared words between consecutive chunks
    """
    if chunk_size <= 0:
        return []
    if overlap < 0:
        overlap = 0
    if overlap >= chunk_size:
        raise ValueError("overlap must be smaller than chunk_size")

    words = text.split()
    chunks = []
    start = 0
    step = chunk_size - overlap
    while start < len(words):
        end = start + chunk_size
        chunk = " ".join(words[start:end])
        chunks.append(chunk)
        start += step
    return chunks
